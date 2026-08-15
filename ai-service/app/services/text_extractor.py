import hashlib
import io
import re
from dataclasses import dataclass
from typing import Any, Literal

import docx
from pypdf import PdfReader

BlockType = Literal["HEADING", "PARAGRAPH", "TABLE", "LIST"]


@dataclass
class ExtractedBlock:
    text: str
    block_type: BlockType
    heading_level: int | None
    locator_type: str
    locator_value: str
    content_hash: str
    source_page_or_index: int

    def to_dict(self) -> dict[str, Any]:
        return {
            "text": self.text,
            "block_type": self.block_type,
            "heading_level": self.heading_level,
            "locator_type": self.locator_type,
            "locator_value": self.locator_value,
            "content_hash": self.content_hash,
            "source_page_or_index": self.source_page_or_index,
        }


class ExtractedChunk:
    def __init__(self, text: str, locator_type: str, locator_value: str, content_hash: str) -> None:
        self.text = text
        self.locator_type = locator_type
        self.locator_value = locator_value
        self.content_hash = content_hash

    def to_dict(self) -> dict[str, Any]:
        return {
            "text": self.text,
            "locator_type": self.locator_type,
            "locator_value": self.locator_value,
            "content_hash": self.content_hash,
        }


def compute_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _is_heading_heuristic(line: str, prev_blank: bool, next_blank: bool) -> tuple[bool, int | None]:
    """Detect heading using text heuristics without requiring font metadata (R-01 mitigation)."""
    stripped = line.strip()
    if not stripped or len(stripped) > 150:
        return False, None

    # Numbered / Named section headings
    numbered_match = re.match(
        r"^(Chương|CHƯƠNG|Phần|PHẦN|Mục|MỤC|Bài|BÀI|Điều|ĐIỀU)\s+(\d+|[IVXLC]+)", stripped
    )
    if numbered_match:
        prefix = numbered_match.group(1).upper()
        level = 1 if prefix in ("CHƯƠNG", "PHẦN") else 2
        return True, level

    # Roman numeral headings: I. Title, II. Title
    if re.match(r"^[IVXLC]+[\.\)]\s+", stripped):
        return True, 1

    # Markdown headings
    md_match = re.match(r"^(#{1,6})\s+(.+)", stripped)
    if md_match:
        return True, len(md_match.group(1))

    # All UPPERCASE short standalone text
    if stripped.isupper() and len(stripped) >= 3 and len(stripped) <= 100:
        return True, 1

    # Short standalone text surrounded by blank lines
    if len(stripped) < 80 and prev_blank and next_blank:
        return True, 2

    return False, None


def extract_pdf_blocks(file_bytes: bytes) -> list[ExtractedBlock]:
    """Extract structured blocks from PDF with heuristic heading detection."""
    reader = PdfReader(io.BytesIO(file_bytes))
    blocks: list[ExtractedBlock] = []

    for i, page in enumerate(reader.pages, start=1):
        raw_text = page.extract_text() or ""
        lines = raw_text.splitlines()

        for idx, line in enumerate(lines):
            stripped = line.strip()
            if not stripped:
                continue

            prev_blank = (idx == 0) or not lines[idx - 1].strip()
            next_blank = (idx == len(lines) - 1) or not lines[idx + 1].strip()

            is_heading, level = _is_heading_heuristic(stripped, prev_blank, next_blank)
            block_type: BlockType = "HEADING" if is_heading else "PARAGRAPH"
            locator_value = f"page:{i}"

            blocks.append(
                ExtractedBlock(
                    text=stripped,
                    block_type=block_type,
                    heading_level=level if is_heading else None,
                    locator_type="PDF_PAGE",
                    locator_value=locator_value,
                    content_hash=compute_hash(stripped),
                    source_page_or_index=i,
                )
            )

    return blocks


def extract_docx_blocks(file_bytes: bytes) -> list[ExtractedBlock]:
    """Extract structured blocks from DOCX using native paragraph styles and table cells."""
    doc = docx.Document(io.BytesIO(file_bytes))
    blocks: list[ExtractedBlock] = []
    block_counter = 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        block_counter += 1
        style_name = para.style.name if para.style else ""
        is_heading = "Heading" in style_name or style_name == "Title"

        heading_level = None
        if is_heading:
            if style_name[-1].isdigit():
                heading_level = int(style_name[-1])
            else:
                heading_level = 1

        block_type: BlockType = "HEADING" if is_heading else "PARAGRAPH"
        locator_value = f"paragraph:{block_counter}"

        blocks.append(
            ExtractedBlock(
                text=text,
                block_type=block_type,
                heading_level=heading_level,
                locator_type="DOCX_PARAGRAPH",
                locator_value=locator_value,
                content_hash=compute_hash(text),
                source_page_or_index=block_counter,
            )
        )

    # Extract tables
    for table_idx, table in enumerate(doc.tables, start=1):
        for row_idx, row in enumerate(table.rows, start=1):
            for cell_idx, cell in enumerate(row.cells, start=1):
                cell_text = cell.text.strip()
                if not cell_text:
                    continue
                block_counter += 1
                locator_value = f"table:{table_idx}/row:{row_idx}/cell:{cell_idx}"
                blocks.append(
                    ExtractedBlock(
                        text=cell_text,
                        block_type="TABLE",
                        heading_level=None,
                        locator_type="DOCX_TABLE_CELL",
                        locator_value=locator_value,
                        content_hash=compute_hash(cell_text),
                        source_page_or_index=block_counter,
                    )
                )

    return blocks


def extract_txt_blocks(file_bytes: bytes) -> list[ExtractedBlock]:
    """Extract structured blocks from TXT using heuristics."""
    text_content = file_bytes.decode("utf-8", errors="replace")
    lines = text_content.splitlines()
    blocks: list[ExtractedBlock] = []

    for idx, line in enumerate(lines, start=1):
        stripped = line.strip()
        if not stripped:
            continue

        prev_blank = (idx == 1) or not lines[idx - 2].strip()
        next_blank = (idx == len(lines)) or not lines[idx].strip()

        is_heading, level = _is_heading_heuristic(stripped, prev_blank, next_blank)
        block_type: BlockType = "HEADING" if is_heading else "PARAGRAPH"
        locator_value = f"line:{idx}-{idx}"

        blocks.append(
            ExtractedBlock(
                text=stripped,
                block_type=block_type,
                heading_level=level if is_heading else None,
                locator_type="TXT_LINE_RANGE",
                locator_value=locator_value,
                content_hash=compute_hash(stripped),
                source_page_or_index=idx,
            )
        )

    return blocks


def extract_blocks(file_bytes: bytes, media_type: str) -> list[ExtractedBlock]:
    """Main block extraction entry point."""
    if media_type == "application/pdf":
        return extract_pdf_blocks(file_bytes)
    elif media_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/docx",
    ):
        return extract_docx_blocks(file_bytes)
    elif media_type == "text/plain":
        return extract_txt_blocks(file_bytes)
    else:
        raise ValueError(f"Định dạng tệp không được hỗ trợ: {media_type}")


# Legacy Adapters
def extract_pdf(file_bytes: bytes) -> list[ExtractedChunk]:
    blocks = extract_pdf_blocks(file_bytes)
    return [
        ExtractedChunk(
            text=b.text,
            locator_type=b.locator_type,
            locator_value=b.locator_value,
            content_hash=b.content_hash,
        )
        for b in blocks
    ]


def extract_docx(file_bytes: bytes) -> list[ExtractedChunk]:
    blocks = extract_docx_blocks(file_bytes)
    return [
        ExtractedChunk(
            text=b.text,
            locator_type=b.locator_type,
            locator_value=b.locator_value,
            content_hash=b.content_hash,
        )
        for b in blocks
    ]


def extract_txt(file_bytes: bytes) -> list[ExtractedChunk]:
    blocks = extract_txt_blocks(file_bytes)
    return [
        ExtractedChunk(
            text=b.text,
            locator_type=b.locator_type,
            locator_value=b.locator_value,
            content_hash=b.content_hash,
        )
        for b in blocks
    ]


def extract_document(file_bytes: bytes, media_type: str) -> list[ExtractedChunk]:
    blocks = extract_blocks(file_bytes, media_type)
    return [
        ExtractedChunk(
            text=b.text,
            locator_type=b.locator_type,
            locator_value=b.locator_value,
            content_hash=b.content_hash,
        )
        for b in blocks
    ]
